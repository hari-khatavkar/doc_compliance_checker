import os
import json
import csv
import argparse
from typing import List, Dict, Any, Optional, Tuple
import PyPDF2
import docx
import tiktoken
import openai
import logging
import datetime
from tqdm import tqdm

# Set up logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

class ComplianceRule:
    """Class to represent a compliance rule."""
    
    def __init__(self, rule_id: str, category: str, short_description: str, full_description: str, active: bool):
        """
        Initialize a compliance rule.
        
        Args:
            rule_id: Unique identifier for the rule
            category: Category of the rule
            short_description: Short description of the rule
            full_description: Full description of the rule
            active: Whether the rule is active
        """
        self.rule_id = rule_id
        self.category = category
        self.short_description = short_description
        self.full_description = full_description
        self.active = active
        
    def __str__(self):
        return f"{self.rule_id}: {self.short_description}"
        
    def to_dict(self):
        """Convert rule to dictionary."""
        return {
            "rule_id": self.rule_id,
            "category": self.category,
            "short_description": self.short_description,
            "full_description": self.full_description,
            "active": self.active
        }

class DocumentComplianceChecker:
    """A class to check document compliance using OpenAI's GPT model."""
    
    def __init__(self, api_key: str, model: str = "gpt-4o-mini"):
        """
        Initialize the Document Compliance Checker.
        
        Args:
            api_key: OpenAI API key
            model: OpenAI model to use (default: gpt-4o-mini)
        """
        self.api_key = api_key
        self.model = model
        openai.api_key = api_key
        self.process_logs = []
        
        # Log initialization
        self.log_process("Initializing Document Compliance Checker", {"model": model})
        
        # GPT-4o-mini context window is 128K tokens
        self.max_tokens = 128000 if "gpt-4" in model else 16385
        # Reserve 1000 tokens for the prompt and 2000 for the response
        self.max_chunk_tokens = self.max_tokens - 3000
        
        # Initialize tokenizer for the model
        self.tokenizer = tiktoken.encoding_for_model(model)
        
    def log_process(self, action: str, details: Dict[str, Any] = None) -> None:
        """
        Log a process step.
        
        Args:
            action: Description of the action
            details: Additional details about the action
        """
        timestamp = datetime.datetime.now().isoformat()
        log_entry = {
            "timestamp": timestamp,
            "action": action,
            "details": details or {}
        }
        self.process_logs.append(log_entry)
        logger.info(f"{action}: {json.dumps(details) if details else ''}")
        
    def extract_text_from_pdf(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text from a PDF file.
        
        Args:
            file_path: Path to the PDF file
            
        Returns:
            Tuple containing:
                - The extracted text
                - A dictionary mapping text to page numbers and positions
        """
        self.log_process(f"Extracting text from PDF", {"file_path": file_path})
        text = ""
        text_mapping = {}
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    page_text = page.extract_text()
                    
                    # Add mapping info (approximate line numbers)
                    lines = page_text.split('\n')
                    start_pos = len(text)
                    for line_num, line in enumerate(lines):
                        end_pos = start_pos + len(line)
                        text_mapping[(start_pos, end_pos)] = {
                            "page": page_num + 1,
                            "line": line_num + 1,
                            "text": line
                        }
                        start_pos = end_pos + 1  # +1 for the newline
                    
                    text += page_text + "\n\n"
                    
            self.log_process("PDF text extraction completed", {
                "pages": len(pdf_reader.pages),
                "total_chars": len(text)
            })
            return text, text_mapping
        except Exception as e:
            error_msg = f"Error extracting text from PDF: {str(e)}"
            self.log_process("PDF text extraction failed", {"error": str(e)})
            logger.error(error_msg)
            raise
            
    def extract_text_from_docx(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text from a DOCX file.
        
        Args:
            file_path: Path to the DOCX file
            
        Returns:
            Tuple containing:
                - The extracted text
                - A dictionary mapping text to page numbers and positions (approximated for DOCX)
        """
        self.log_process("Extracting text from DOCX", {"file_path": file_path})
        text = ""
        text_mapping = {}
        
        try:
            doc = docx.Document(file_path)
            
            # For DOCX, we don't have direct page mapping, so we'll approximate
            # based on paragraphs (assuming 40 paragraphs per page)
            paragraphs_per_page = 40
            current_pos = 0
            
            for i, para in enumerate(doc.paragraphs):
                para_text = para.text
                if not para_text.strip():
                    continue
                
                # Calculate approximate page and line
                approx_page = (i // paragraphs_per_page) + 1
                approx_line = (i % paragraphs_per_page) + 1
                
                # Add mapping
                text_mapping[(current_pos, current_pos + len(para_text))] = {
                    "page": approx_page,
                    "line": approx_line,
                    "text": para_text
                }
                
                text += para_text + "\n"
                current_pos += len(para_text) + 1  # +1 for the newline
                
            self.log_process("DOCX text extraction completed", {
                "paragraphs": len(doc.paragraphs),
                "total_chars": len(text)
            })
            return text, text_mapping
        except Exception as e:
            error_msg = f"Error extracting text from DOCX: {str(e)}"
            self.log_process("DOCX text extraction failed", {"error": str(e)})
            logger.error(error_msg)
            raise
            
    def extract_text(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text from a file based on its extension.
        
        Args:
            file_path: Path to the file
            
        Returns:
            Tuple containing:
                - The extracted text
                - A dictionary mapping text to locations in the document
        """
        file_extension = os.path.splitext(file_path)[1].lower()
        
        if file_extension == '.pdf':
            return self.extract_text_from_pdf(file_path)
        elif file_extension == '.docx':
            return self.extract_text_from_docx(file_path)
        else:
            error_msg = f"Unsupported file type: {file_extension}. Only PDF and DOCX are supported."
            self.log_process("File type not supported", {"file_extension": file_extension})
            raise ValueError(error_msg)
            
    def split_text_into_chunks(self, text: str) -> List[Tuple[str, int, int]]:
        """
        Split text into chunks that fit within the model's context window.
        
        Args:
            text: The text to split
            
        Returns:
            A list of tuples containing (chunk_text, start_pos, end_pos)
        """
        self.log_process("Splitting text into chunks")
        tokens = self.tokenizer.encode(text)
        chunks = []
        current_chunk = []
        current_chunk_length = 0
        chunk_start_pos = 0
        
        for i, token in enumerate(tokens):
            if current_chunk_length + 1 <= self.max_chunk_tokens:
                current_chunk.append(token)
                current_chunk_length += 1
            else:
                # Calculate the position in the original text
                chunk_text = self.tokenizer.decode(current_chunk)
                chunk_end_pos = chunk_start_pos + len(chunk_text)
                chunks.append((chunk_text, chunk_start_pos, chunk_end_pos))
                
                # Start a new chunk
                current_chunk = [token]
                current_chunk_length = 1
                chunk_start_pos = chunk_end_pos
                
        if current_chunk:
            chunk_text = self.tokenizer.decode(current_chunk)
            chunk_end_pos = chunk_start_pos + len(chunk_text)
            chunks.append((chunk_text, chunk_start_pos, chunk_end_pos))
            
        self.log_process("Text chunking completed", {"chunks": len(chunks)})
        return chunks
        
    def analyze_chunk(self, chunk_data: Tuple[str, int, int], chunk_number: int, total_chunks: int, 
                      compliance_rules: List[ComplianceRule]) -> Dict[str, Any]:
        """
        Analyze a chunk of text for compliance using GPT.
        
        Args:
            chunk_data: Tuple of (chunk_text, start_pos, end_pos)
            chunk_number: The current chunk number
            total_chunks: The total number of chunks
            compliance_rules: The compliance rules to check against
            
        Returns:
            The compliance analysis results for the chunk
        """
        chunk_text, start_pos, end_pos = chunk_data
        self.log_process(f"Analyzing chunk {chunk_number}/{total_chunks}", {
            "chunk_size": len(chunk_text),
            "start_pos": start_pos,
            "end_pos": end_pos
        })
        
        # Format rules for prompt
        rules_text = ""
        for rule in compliance_rules:
            if rule.active:
                rules_text += f"Rule ID: {rule.rule_id}\n"
                rules_text += f"Category: {rule.category}\n"
                rules_text += f"Description: {rule.short_description}\n"
                rules_text += f"Details: {rule.full_description}\n\n"
        
        prompt = f"""
You are an AI document compliance checker. Analyze the following document text for compliance with the rules listed below.
This is chunk {chunk_number} of {total_chunks}.

COMPLIANCE RULES:
{rules_text}

YOUR TASK:
1. Analyze the document text for compliance with each active rule.
2. For each rule, determine if the document is compliant, non-compliant, or if compliance cannot be determined from this chunk.
3. For non-compliant items, provide specific examples from the text and explain why they violate the rule.
4. If possible, identify the approximate location in the text where the violation occurs (character position in the chunk).
5. Provide recommendations for resolving each non-compliance issue.
6. Format your response as a JSON object with the following structure:
   {{
     "chunk_number": {chunk_number},
     "chunk_start_pos": {start_pos},
     "chunk_end_pos": {end_pos},
     "results": [
       {{
         "rule_id": "Rule ID",
         "status": "compliant" | "non_compliant" | "undetermined",
         "evidence": "Evidence or examples from the text (only for non-compliant items)",
         "explanation": "Explanation of why it violates the rule (only for non-compliant items)",
         "position_in_chunk": "Approximate character position in the chunk where violation occurs",
         "recommendation": "Recommendation to fix/resolve the issue"
       }},
       ...
     ]
   }}

DOCUMENT TEXT:
{chunk_text}

Respond with only the JSON object. Do not include any other text in your response.
"""

        try:
            response = openai.chat.completions.create(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                temperature=0.0,
                response_format={"type": "json_object"}
            )
            
            analysis = json.loads(response.choices[0].message.content)
            self.log_process(f"Chunk {chunk_number} analysis completed", {
                "rules_analyzed": len(analysis.get("results", [])),
                "token_usage": response.usage.total_tokens if hasattr(response, 'usage') else "unknown"
            })
            return analysis
        except Exception as e:
            error_msg = f"Error analyzing chunk {chunk_number}: {str(e)}"
            self.log_process(f"Chunk {chunk_number} analysis failed", {"error": str(e)})
            logger.error(error_msg)
            return {
                "chunk_number": chunk_number,
                "chunk_start_pos": start_pos,
                "chunk_end_pos": end_pos,
                "results": [],
                "error": str(e)
            }
    
    def load_compliance_rules(self, rules_file: str) -> List[ComplianceRule]:
        """
        Load compliance rules from a CSV file.
        
        Args:
            rules_file: Path to the CSV file containing compliance rules
            
        Returns:
            List of ComplianceRule objects
        """
        self.log_process("Loading compliance rules", {"file": rules_file})
        rules = []
        
        try:
            with open(rules_file, 'r', newline='', encoding='utf-8') as csvfile:
                reader = csv.DictReader(csvfile)
                for row in reader:
                    # Check for required columns
                    required_cols = ["Rule ID", "Category", "Short Description", "Full Description", "Active"]
                    if not all(col in row for col in required_cols):
                        missing = [col for col in required_cols if col not in row]
                        raise ValueError(f"Missing columns in CSV: {', '.join(missing)}")
                    
                    # Convert Active field to boolean
                    active = row["Active"].lower() in ["true", "yes", "1", "active"]
                    
                    rule = ComplianceRule(
                        rule_id=row["Rule ID"],
                        category=row["Category"],
                        short_description=row["Short Description"],
                        full_description=row["Full Description"],
                        active=active
                    )
                    rules.append(rule)
                    
            self.log_process("Compliance rules loaded", {
                "total_rules": len(rules),
                "active_rules": sum(1 for rule in rules if rule.active)
            })
            return rules
        except Exception as e:
            error_msg = f"Error loading compliance rules: {str(e)}"
            self.log_process("Compliance rules loading failed", {"error": str(e)})
            logger.error(error_msg)
            raise
            
    def analyze_document(self, file_path: str, rules_file: str) -> Dict[str, Any]:
        """
        Analyze a document for compliance.
        
        Args:
            file_path: Path to the document
            rules_file: Path to the CSV file containing compliance rules
            
        Returns:
            The compliance analysis report
        """
        file_id = os.path.basename(file_path)
        file_name = os.path.splitext(file_id)[0]
        self.log_process(f"Starting compliance analysis", {"file_id": file_id, "rules_file": rules_file})
        
        # Load compliance rules
        compliance_rules = self.load_compliance_rules(rules_file)
        active_rules = [rule for rule in compliance_rules if rule.active]
        
        if not active_rules:
            self.log_process("No active compliance rules found")
            return {
                "file_id": file_id,
                "file_name": file_name,
                "timestamp": datetime.datetime.now().isoformat(),
                "overall_status": "compliant",
                "message": "No active compliance rules to check against",
                "rule_violations": []
            }
        
        # Extract text from document
        text, text_mapping = self.extract_text(file_path)
        
        # Split text into chunks
        chunks = self.split_text_into_chunks(text)
        
        # Analyze each chunk
        chunk_results = []
        for i, chunk_data in enumerate(tqdm(chunks, desc="Analyzing chunks")):
            result = self.analyze_chunk(chunk_data, i+1, len(chunks), active_rules)
            chunk_results.append(result)
            
        # Compile the final report
        report = self.compile_report(file_id, file_name, active_rules, chunk_results, text_mapping)
        return report
        
    def find_location_in_document(self, absolute_pos: int, text_mapping: Dict) -> Dict[str, Any]:
        """
        Find the location (page, line) in the document based on absolute position.
        
        Args:
            absolute_pos: Absolute position in the document text
            text_mapping: Mapping from text positions to document locations
            
        Returns:
            Dictionary with page and line information
        """
        for (start, end), location in text_mapping.items():
            if start <= absolute_pos < end:
                return location
                
        # If no exact match, find the closest
        closest_start = None
        closest_location = None
        min_distance = float('inf')
        
        for (start, end), location in text_mapping.items():
            distance = min(abs(start - absolute_pos), abs(end - absolute_pos))
            if distance < min_distance:
                min_distance = distance
                closest_start = start
                closest_location = location
                
        if closest_location:
            return closest_location
            
        return {"page": "unknown", "line": "unknown", "text": ""}
        
    def compile_report(self, file_id: str, file_name: str, compliance_rules: List[ComplianceRule], 
                       chunk_results: List[Dict[str, Any]], text_mapping: Dict) -> Dict[str, Any]:
        """
        Compile the final compliance report from chunk results.
        
        Args:
            file_id: ID of the document (filename)
            file_name: Name of the document
            compliance_rules: List of compliance rules that were checked
            chunk_results: Results from analyzing each chunk
            text_mapping: Mapping from text positions to document locations
            
        Returns:
            The compiled compliance report
        """
        self.log_process("Compiling final compliance report")
        
        # Create a map of rule_id to rule details
        rule_map = {rule.rule_id: rule for rule in compliance_rules}
        
        # Initialize structure to hold compliance results
        compliance_status = {}
        violations = {}
        
        for rule in compliance_rules:
            compliance_status[rule.rule_id] = "compliant"
            violations[rule.rule_id] = []
            
        # Process chunk results
        for chunk_result in chunk_results:
            if "error" in chunk_result:
                continue
                
            chunk_start_pos = chunk_result.get("chunk_start_pos", 0)
            
            for result in chunk_result.get("results", []):
                rule_id = result.get("rule_id")
                status = result.get("status")
                
                if rule_id not in compliance_status:
                    continue
                    
                # If any chunk shows non-compliance, the rule is non-compliant
                if status == "non_compliant":
                    compliance_status[rule_id] = "non_compliant"
                    
                    # Calculate absolute position in document
                    position_in_chunk = result.get("position_in_chunk", "0")
                    try:
                        pos_in_chunk = int(position_in_chunk)
                    except (ValueError, TypeError):
                        pos_in_chunk = 0
                    
                    absolute_pos = chunk_start_pos + pos_in_chunk
                    location = self.find_location_in_document(absolute_pos, text_mapping)
                    
                    violation = {
                        "evidence": result.get("evidence", ""),
                        "explanation": result.get("explanation", ""),
                        "location": {
                            "page": location.get("page", "unknown"),
                            "line": location.get("line", "unknown"),
                            "text_context": location.get("text", "")
                        },
                        "recommendation": result.get("recommendation", "")
                    }
                    
                    violations[rule_id].append(violation)
                    
        # Determine overall compliance status
        overall_status = "compliant"
        for status in compliance_status.values():
            if status == "non_compliant":
                overall_status = "non_compliant"
                break
                
        # Compile list of violations for the report
        rule_violations = []
        for rule_id, status in compliance_status.items():
            if status == "non_compliant":
                rule = rule_map.get(rule_id)
                if not rule:
                    continue
                    
                rule_violation = {
                    "rule_id": rule_id,
                    "category": rule.category,
                    "short_description": rule.short_description,
                    "violations": violations[rule_id]
                }
                
                rule_violations.append(rule_violation)
                
        # Create final report
        timestamp = datetime.datetime.now().isoformat()
        report = {
            "file_id": file_id,
            "file_name": file_name,
            "overall_compliance_status": overall_status,
            "timestamp": timestamp,
            "summary": {
                "total_rules": len(compliance_rules),
                "compliant_rules": sum(1 for status in compliance_status.values() if status == "compliant"),
                "non_compliant_rules": sum(1 for status in compliance_status.values() if status == "non_compliant")
            },
            "rule_violations": rule_violations
        }
        
        self.log_process("Final report compiled", {
            "overall_status": overall_status,
            "total_violations": len(rule_violations)
        })
        
        return report
        
    def save_report(self, report: Dict[str, Any], output_path: str) -> None:
        """
        Save the compliance report to a file.
        
        Args:
            report: The compliance report
            output_path: Path to save the report
        """
        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(report, f, indent=2)
            self.log_process("Report saved", {"output_path": output_path})
        except Exception as e:
            error_msg = f"Error saving report: {str(e)}"
            self.log_process("Error saving report", {"error": str(e)})
            logger.error(error_msg)
            raise
            
    def save_process_logs(self, output_path: str) -> None:
        """
        Save the process logs to a CSV file.
        
        Args:
            output_path: Path to save the process logs
        """
        try:
            with open(output_path, 'w', newline='', encoding='utf-8') as csvfile:
                fieldnames = ['timestamp', 'action', 'details']
                writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
                writer.writeheader()
                
                for log in self.process_logs:
                    writer.writerow({
                        'timestamp': log['timestamp'],
                        'action': log['action'],
                        'details': json.dumps(log['details']) if log['details'] else ''
                    })
                    
            logger.info(f"Process logs saved to {output_path}")
        except Exception as e:
            logger.error(f"Error saving process logs: {e}")
            raise

def main():
    """Main function to run the document compliance checker."""
    parser = argparse.ArgumentParser(description='AI-Based Document Compliance Checker')
    parser.add_argument('--file', required=True, help='Path to the document file (PDF or DOCX)')
    parser.add_argument('--rules', default='compliance_rules.csv', 
                        help='Path to CSV file containing compliance rules (default: compliance_rules.csv)')
    parser.add_argument('--output', default='compliance_report.json', 
                        help='Path to save the compliance report (default: compliance_report.json)')
    parser.add_argument('--log-output', default='process_logs.csv',
                        help='Path to save the process logs (default: process_logs.csv)')
    parser.add_argument('--api-key', required=True, help='OpenAI API key')
    parser.add_argument('--model', default='gpt-4o-mini', help='OpenAI model to use (default: gpt-4o-mini)')
    
    args = parser.parse_args()
    
    try:
        # Initialize checker
        checker = DocumentComplianceChecker(api_key=args.api_key, model=args.model)
        
        # Analyze document
        report = checker.analyze_document(args.file, args.rules)
        
        # Save report
        checker.save_report(report, args.output)
        
        # Save process logs
        checker.save_process_logs(args.log_output)
        
        logger.info("Compliance analysis completed successfully")
        
    except Exception as e:
        logger.error(f"Error in compliance analysis: {e}", exc_info=True)
        return 1
        
    return 0

if __name__ == "__main__":
    exit(main())